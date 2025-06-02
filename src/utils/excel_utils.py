"""
Provides utility functions for working with Excel files, focusing on styling and formatting.
"""
import logging

from openpyxl.styles import Font, Border, Side, Alignment, PatternFill
from openpyxl.utils import get_column_letter


logger = logging.getLogger(__name__)


def apply_default_report_styles(
    ws,
    num_header_rows=1,
    num_footer_rows=1,
    center_columns=None,         # 예: ['인원']
    number_columns=None,         # 예: ['금액', '합계']
    # 고급: {'합계': {'align': ..., 'format': ...}, ...}
    column_style_map=None
):
    """
    헤더명/컬럼명 기반으로 스타일을 유연하게 적용하는 개선된 함수.
    center_columns: 가운데 정렬할 컬럼명 리스트
    number_columns: 숫자 포맷 적용할 컬럼명 리스트
    column_style_map: 컬럼별 고급 스타일 매핑 dict (선택)
    """
    font_format = Font(size=11, name='맑은 고딕')
    font_format_bold = Font(size=11, name='맑은 고딕', bold=True)
    border_format = Side(border_style="thin")
    align_format_center = Alignment(horizontal="center", vertical="center")
    align_format_left = Alignment(horizontal="left", vertical="center")
    fill_style = PatternFill(start_color="00C0C0C0",
                             end_color="00C0C0C0", patternType="solid")

    # 헤더 추출
    header_row = [cell.value for cell in ws[1]]

    # 옵션 기본값
    center_columns = center_columns or []
    number_columns = number_columns or []
    column_style_map = column_style_map or {}

    max_row = ws.max_row

    for row_idx, row_cells in enumerate(ws.iter_rows()):
        for col_idx, cell in enumerate(row_cells):
            col_name = header_row[col_idx] if col_idx < len(
                header_row) else None

            # 기본 스타일
            cell.font = font_format
            cell.border = Border(
                top=border_format, bottom=border_format, left=border_format, right=border_format)
            cell.alignment = align_format_left

            # 헤더
            if row_idx < num_header_rows:
                cell.font = font_format_bold
                cell.fill = fill_style
                cell.alignment = align_format_center

            # 푸터(합계 등)
            elif row_idx >= (max_row - num_footer_rows):
                cell.font = font_format_bold
                cell.fill = fill_style
                # 호환성을 위해 col_idx<2도 유지
                if col_name and (col_name in center_columns or col_idx < 2):
                    cell.alignment = align_format_center

            # 데이터 행
            else:
                # 고급 스타일 맵 우선 적용
                style = column_style_map.get(col_name, {})
                if style.get('align'):
                    cell.alignment = style['align']
                elif col_name in center_columns:
                    cell.alignment = align_format_center

                # 숫자 포맷 적용
                if style.get('format') and isinstance(cell.value, (int, float)):
                    cell.number_format = style['format']
                elif col_name in number_columns and isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0'

    # 컬럼 폭 자동 조정
    for col_idx, column_cells in enumerate(ws.columns):
        try:
            relevant_values = [str(cell.value)
                               for cell in column_cells if cell.value is not None]
            if not relevant_values:
                new_column_length = 10
            else:
                new_column_length = max(len(value)
                                        for value in relevant_values)
            new_column_letter = get_column_letter(column_cells[0].column)
            adjusted_width = (new_column_length * 1.2) + 2
            if adjusted_width < 10:
                adjusted_width = 10
            ws.column_dimensions[new_column_letter].width = adjusted_width
        except Exception as e:
            logger.warning(f"열 자동 너비 설정 실패 (열 {col_idx}): {e}")
